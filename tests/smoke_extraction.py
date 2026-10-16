"""Real-model smoke test. Run with venv/bin/python tests/smoke_extraction.py."""
import argparse
import json
from pathlib import Path
import tempfile
import time

import requests
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook

INVOICE = [
    "SAMPLE SUPPLIES", "Kathmandu, Nepal", "INVOICE", "Bill No: INV-001",
    "Date: 2026-10-16", "Bill To: Test Customer", "Payment Mode: Cash", "",
    "Sno  Particular   Quantity   Rate   Amount",
    "1    Notebook     2 pcs      50     100", "",
    "Sub Total: 100.00", "Discount: 0.00", "Taxable Amount: 100.00",
    "Tax: 13.00", "Grand Total: 113.00",
    "Amount in words: One Hundred Thirteen Only",
]


def create_samples(directory):
    image = Image.new("RGB", (1400, 1200), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("DejaVuSansMono.ttf", 32)
    for index, line in enumerate(INVOICE):
        draw.text((60, 40 + index * 60), line, font=font, fill="black")
    image.save(directory / "invoice.png")
    image.save(directory / "invoice.pdf", "PDF", resolution=120)
    doc = Document()
    for line in INVOICE:
        doc.add_paragraph(line)
    doc.save(directory / "invoice.docx")
    workbook = Workbook()
    for line in INVOICE:
        workbook.active.append([line])
    workbook.save(directory / "invoice.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:8000")
    parser.add_argument("--formats", nargs="+", default=["png", "pdf", "docx", "xlsx"])
    parser.add_argument("--output", type=Path, default=Path("/tmp/invoice-smoke-results.json"))
    args = parser.parse_args()
    for attempt in range(30):
        try:
            health = requests.get(f"{args.url}/health", timeout=5)
            health.raise_for_status()
            if health.json().get("status") == "healthy":
                break
        except requests.RequestException:
            pass
        time.sleep(1)
    else:
        raise RuntimeError("API is not ready; check /health and configured models")
    results = []
    with tempfile.TemporaryDirectory() as temp:
        directory = Path(temp)
        create_samples(directory)
        for extension in args.formats:
            path = directory / f"invoice.{extension}"
            started = time.monotonic()
            with path.open("rb") as file:
                response = requests.post(f"{args.url}/extract", params={"include_raw": "true"},
                                         files={"file": (path.name, file)}, timeout=900)
            response.raise_for_status()
            result = response.json()
            results.append({"format": extension, "seconds": round(time.monotonic()-started, 2), "response": result})
            args.output.write_text(json.dumps(results, indent=2))
            assert result["success"], result
            data = result["data"]
            assert result["stage1_output"], result
            assert "INV-001" in data["bill_no"], data
            assert data["items"], data
            assert abs(data["grand_total"] - 113) < 0.01, data
            print(f"PASS {extension}: bill={data['bill_no']}, total={data['grand_total']}, {results[-1]['seconds']}s", flush=True)
    print(f"Saved results: {args.output}")


if __name__ == "__main__":
    main()
